#!/usr/bin/env python3
"""Conservative reverse cross-release overlap screen.

Compare BraTS 2021 training/validation processed 2D image slices to BraTS 2020
held-out TEST full volumes. Exact matches of >=3 DISTINCT informative slice hashes
for the SAME patient pair are strong SCREENING CANDIDATES, NOT verified identity.

Outputs a conservative candidate-exclusion file for the companion evaluation:
    brats2020_exclusion_candidates.txt
The file is a sensitivity-analysis list, not a list of confirmed duplicates.

Requirements: numpy; optional openpyxl for highlighted workbook.
"""
import argparse
import csv
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
import numpy as np

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    Workbook = None

IMAGE_KEYS = ('image', 'images', 'image_cropped', 'volume')


def args_parser():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--train_images21', default='processed_data21/trainImage')
    p.add_argument('--val_images21', default='processed_data21/valImage')
    p.add_argument('--test_volumes20', default='processed_data/testVolumes')
    p.add_argument('--manifest20', default='processed_data/manifests/patient_split.json')
    p.add_argument('--manifest21', default='processed_data21/manifests/patient_split.json')
    p.add_argument('--output_dir', default='brats2021_to_brats2020_reverse_overlap_audit')
    p.add_argument('--min_matching_slices', type=int, default=3)
    p.add_argument('--min_nonzero_fraction', type=float, default=.005)
    p.add_argument('--min_spatial_std', type=float, default=1e-4)
    p.add_argument('--allow_incomplete_source_coverage', action='store_true',
                   help='Continue if some 2021 training/validation patients have no processed slices; document reduced coverage.')
    return p.parse_args()


def read_splits(path):
    with open(path, encoding='utf-8') as fh:
        j = json.load(fh)
    keys = ('train_patients', 'validation_patients', 'test_patients')
    splits = {}
    for k in keys:
        ids = j.get(k)
        if not isinstance(ids, list) or not ids or len(ids) != len(set(ids)):
            raise ValueError(f'{path}: missing/invalid split {k}')
        splits[k] = set(map(str, ids))
    if any(splits[a] & splits[b] for i,a in enumerate(keys) for b in keys[i+1:]):
        raise ValueError(f'{path}: within-release patient overlap')
    return splits


def load_img(path):
    with np.load(path, allow_pickle=False) as d:
        key = next((k for k in IMAGE_KEYS if k in d), None)
        if key is None:
            raise ValueError(f'{path}: no image key')
        a = np.asarray(d[key])
    if a.ndim != 4 or a.shape[-1] != 4:
        raise ValueError(f'{path}: expected [Z,H,W,4], got {a.shape}')
    return a


def informative(a, fraction, spatial_std):
    if a.ndim != 3 or a.shape[-1] != 4 or not np.isfinite(a).all():
        return False
    return (np.any(np.abs(a) > 1e-8, axis=-1).mean() >= fraction
            and np.std(a.astype(np.float32, copy=False), axis=(0,1)).mean() >= spatial_std)


def sha(a):
    a = np.ascontiguousarray(np.asarray(a, dtype=np.float32))
    h = hashlib.sha256()
    h.update(str(a.shape).encode())
    h.update(memoryview(a).cast('B'))
    return h.hexdigest()


def write_csv(path, rows, columns):
    with open(path, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=columns)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, summary, pairs, patients):
    if Workbook is None:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws.append(['Measure', 'Value'])
    for k,v in summary.items():
        ws.append([k, json.dumps(v) if isinstance(v,(dict,list)) else v])
    for name, rows, columns in (
        ('Patient pair candidates', pairs, ['status','brats2020_patient','brats2021_patient','brats2021_split','distinct_informative_slices20','distinct_exact_hashes']),
        ('BraTS2020 patient decisions', patients, ['brats2020_patient','screening_status','matched_2021_patient_ids','candidate_pairs'])):
        ws = wb.create_sheet(name)
        ws.append(columns)
        for row in rows:
            ws.append([row.get(c, '') for c in columns])
            if row.get('status', row.get('screening_status', '')).startswith('STRONG'):
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill('solid', fgColor='FFC7CE')
                    cell.font = Font(color='9C0006', bold=True)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    a = args_parser()
    if a.min_matching_slices < 2:
        raise ValueError('Minimum matching informative slices must be at least 2; default is 3')
    out = Path(a.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    split20, split21 = read_splits(a.manifest20), read_splits(a.manifest21)
    folders = {'train': Path(a.train_images21), 'validation': Path(a.val_images21)}
    volume_folder = Path(a.test_volumes20)
    for folder in [*folders.values(), volume_folder]:
        if not folder.is_dir():
            raise FileNotFoundError(f'Required folder not found: {folder}')

    # Map 2021 slices to authoritative patient IDs from the 2021 manifest.
    index = defaultdict(set)
    scanned = {}
    missing_source_patients = {}
    for split, folder in folders.items():
        ids = split21['train_patients' if split == 'train' else 'validation_patients']
        count, useful, invalid = 0, 0, 0
        observed_patients, seen_slice_indices = set(), set()
        files = sorted(folder.glob('*.npy'))
        if not files:
            raise ValueError(f'No NPY slices in {folder}')
        for f in files:
            count += 1
            # Patient names may contain underscores. Use the complete stem minus
            # the LAST numeric slice suffix, and verify against manifest.
            m = re.match(r'^(.*)_(\d+)$', f.stem)
            if not m or m.group(1) not in ids:
                invalid += 1
                continue
            slice_key = (m.group(1), int(m.group(2)))
            if slice_key in seen_slice_indices:
                raise ValueError(f'Duplicate processed slice patient/index: {slice_key}')
            seen_slice_indices.add(slice_key)
            observed_patients.add(m.group(1))
            img = np.load(f, allow_pickle=False)
            if img.ndim != 3 or img.shape[-1] != 4 or not np.isfinite(img).all():
                raise ValueError(f'Invalid source image shape/nonfinite values: {f}')
            if not informative(img, a.min_nonzero_fraction, a.min_spatial_std):
                continue
            useful += 1
            index[sha(img)].add((split, m.group(1)))
        missing = sorted(ids - observed_patients)
        missing_source_patients[split] = missing
        scanned[split] = {'files':count, 'informative':useful, 'unrecognized_id':invalid,
                          'manifest_patient_count':len(ids), 'patients_with_slices':len(observed_patients),
                          'missing_manifest_patients':len(missing)}
        if invalid:
            raise ValueError(f'{folder}: {invalid} slice filenames do not match the corresponding 2021 manifest; verify preprocessing naming')
        if missing and not a.allow_incomplete_source_coverage:
            raise ValueError(
                f'{folder}: {len(missing)} manifest patients have no processed slices; '
                f'first IDs: {missing[:10]}. Audit would have incomplete patient coverage. '
                'Investigate or explicitly use --allow_incomplete_source_coverage.'
            )
        if missing:
            print(f'WARNING: {len(missing)} {split} patients have no processed slices')
        print(f'Indexed 2021 {split}: {useful} informative slices from {count} files')

    volfiles = sorted(volume_folder.glob('*.npz'))
    if not volfiles:
        raise ValueError(f'No NPZ patient volumes in {volume_folder}')
    testids = split20['test_patients']
    found = set(p.stem for p in volfiles)
    if found != testids:
        raise ValueError(f'2020 test files/manifest disagree: missing={sorted(testids-found)[:10]}, unexpected={sorted(found-testids)[:10]}')

    hits = defaultdict(lambda: {'z20':set(), 'hashes':set()})
    target_useful = 0
    for i,f in enumerate(volfiles, 1):
        vol = load_img(f)
        for z in range(vol.shape[0]):
            sl = vol[z]
            if not informative(sl, a.min_nonzero_fraction, a.min_spatial_std):
                continue
            target_useful += 1
            digest = sha(sl)
            for split,pid21 in index.get(digest, ()):
                item = hits[(f.stem, split, pid21)]
                item['z20'].add(z)
                item['hashes'].add(digest)
        if i % 25 == 0:
            print(f'Scanned {i}/{len(volfiles)} BraTS 2020 test patients')

    pairs, strong = [], defaultdict(list)
    for (pid20, split, pid21), h in sorted(hits.items()):
        nz, nh = len(h['z20']), len(h['hashes'])
        status = ('STRONG_CANDIDATE_REVIEW' if min(nz,nh) >= a.min_matching_slices
                  else 'WEAK_CANDIDATE_REVIEW')
        row = {'status':status, 'brats2020_patient':pid20, 'brats2021_patient':pid21,
               'brats2021_split':split, 'distinct_informative_slices20':nz,
               'distinct_exact_hashes':nh}
        pairs.append(row)
        if status.startswith('STRONG'):
            strong[pid20].append(row)
    patients = []
    for pid in sorted(testids):
        candidate = strong.get(pid, [])
        patients.append({'brats2020_patient':pid,
                         'screening_status':'STRONG_CANDIDATE_REVIEW' if candidate else 'NO_STRONG_MATCH_NOT_PROVEN_INDEPENDENT',
                         'matched_2021_patient_ids':';'.join(sorted({r['brats2021_patient'] for r in candidate})),
                         'candidate_pairs':len(candidate)})

    # The companion evaluator reads precisely these IDs. The exclusions are
    # CONSERVATIVE SCREENING EXCLUSIONS, not confirmed subject-identity findings.
    (out/'brats2020_exclusion_candidates.txt').write_text(
        '\n'.join(sorted(strong)) + ('\n' if strong else ''), encoding='utf-8')
    write_csv(out/'reverse_patient_pair_candidates.csv', pairs,
              ['status','brats2020_patient','brats2021_patient','brats2021_split','distinct_informative_slices20','distinct_exact_hashes'])
    write_csv(out/'reverse_patient_screening.csv', patients,
              ['brats2020_patient','screening_status','matched_2021_patient_ids','candidate_pairs'])
    summary = {'direction':'2021 train/validation -> 2020 test',
               'source_patient_coverage_complete':not any(missing_source_patients.values()),
               'missing_2021_source_patient_ids':missing_source_patients,
               'source_coverage_override':a.allow_incomplete_source_coverage,
               '2021_source_slice_scan':scanned, '2020_test_patients_scanned':len(volfiles),
               '2020_informative_slices_scanned':target_useful,
               '2020_strong_candidate_patients':len(strong),
               '2020_patients_without_strong_match_not_verified_independent':len(testids)-len(strong),
               'strong_pair_count':sum(len(v) for v in strong.values()),
               'weak_pair_count':sum(x['status'].startswith('WEAK') for x in pairs),
               'min_distinct_informative_matching_slices':a.min_matching_slices,
               'limitations':['Not a proof of subject identity.',
                              'Exact array matching can miss overlapping patients after different preprocessing.',
                              'Candidate exclusions support a sensitivity analysis only.',
                              'Verify checkpoint train/validation provenance independently.']}
    (out/'reverse_audit_summary.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
    write_xlsx(out/'reverse_overlap_screen_highlighted.xlsx', summary, pairs, patients)
    print('Reverse overlap screen completed. Strong-candidate 2020 patients:',len(strong))
    print('Exclusion file:', out/'brats2020_exclusion_candidates.txt')
    print('Matches are SCREENING flags; no-match does not prove independence.')


if __name__ == '__main__':
    main()
